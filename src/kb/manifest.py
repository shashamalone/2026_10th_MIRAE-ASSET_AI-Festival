# -*- coding: utf-8 -*-
"""2026-08-24 주최측 XLSX 정본 선언과 읽기 전용 검증.

원본은 수정하지 않는다. ``--check`` 빌더와 실제 적재기가 이 모듈을 함께
사용해 파일 집합, 행 수, 헤더, 공식 타입, PK와 관측 기준일 계약이 서로
달라지는 것을 막는다.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATASET_VERSION = "financial-products-2026-08-24"
RELEASE_DATE = date(2026, 8, 24)
EXTERNAL_CUTOFF = RELEASE_DATE
DEFAULT_DATASET_DIR = (
    ROOT.parent / "data" / "ai-festival2026_금융상품Agent_DtataSet260824"
)

SCHEMA_HEADER = ("순번", "컬럼명", "데이터타입", "Nullable", "컬럼코멘트")
TYPE_PATTERN = re.compile(
    r"^(?:text|bigint|double precision|numeric\(\d{1,2},\d{1,2}\))$",
    re.IGNORECASE,
)


class LookAheadError(ValueError):
    """허용 cutoff 뒤의 관측값이 정본이나 외부 근거에 섞였을 때 발생한다."""


@dataclass(frozen=True)
class SourceSpec:
    code: str
    raw_table: str
    data_file: str
    schema_file: str
    expected_rows: int
    expected_columns: int
    primary_key: tuple[str, ...]
    effective_as_of_columns: tuple[str, ...]
    grain: str


SOURCES: tuple[SourceSpec, ...] = (
    SourceSpec(
        "PRBD01N001",
        "bond_kr_master",
        "prbd01n001_data.xlsx",
        "prbd01n001_schema.xlsx",
        21_882,
        58,
        ("pd_no", "pd_exg_mkt", "info_base_dt", "info_seq"),
        ("info_base_dt", "pd_std_info_update", "sale_yield_base_dt"),
        "채권×시장×정보기준일×판매 LOT",
    ),
    SourceSpec(
        "PREF01N001",
        "etf_kr_master",
        "pref01n001_data.xlsx",
        "pref01n001_schema.xlsx",
        1_780,
        98,
        ("pd_itm_no",),
        ("cu_upt_dt", "du_upt_dt", "wu_upt_dt", "fn_base_dt", "ref_base_dt"),
        "국내 ETF/ETN 상품",
    ),
    SourceSpec(
        "PREF02N001",
        "etf_gl_master",
        "pref02n001_data.xlsx",
        "pref02n001_schema.xlsx",
        6_037,
        49,
        ("pd_itm_no",),
        ("cu_upt_dt", "du_upt_dt", "wu_upt_dt", "du_clpr_base_dt", "du_nav_base_dt"),
        "해외 ETF/ETN 상품",
    ),
    SourceSpec(
        "PRFD01N001",
        "fund_pub_master",
        "prfd01n001_data.xlsx",
        "prfd01n001_schema.xlsx",
        23_676,
        75,
        ("itm_no",),
        ("fd_daily_bas_dt", "fd_price_bas_dt"),
        "펀드 상품(공모·사모)",
    ),
)


@dataclass(frozen=True)
class OfficialColumn:
    ordinal: int
    name: str
    data_type: str
    nullable: bool
    description: str


@dataclass(frozen=True)
class SourceInspection:
    spec: SourceSpec
    data_path: Path
    schema_path: Path
    columns: tuple[OfficialColumn, ...]
    null_counts: tuple[int, ...]
    row_count: int
    effective_as_of: date | None
    data_sha256: str
    schema_sha256: str
    excluded_rows: int = 0

    def as_dict(self) -> dict[str, object]:
        return {
            "code": self.spec.code,
            "raw_table": f"raw.{self.spec.raw_table}",
            "data_file": self.data_path.name,
            "schema_file": self.schema_path.name,
            "rows": self.row_count,
            "loaded_rows": self.row_count,
            "excluded_rows": 0,
            "columns": len(self.columns),
            "primary_key": list(self.spec.primary_key),
            "official_nullable_conflicts": {
                column.name: self.null_counts[column.ordinal - 1]
                for column in self.columns
                if not column.nullable and self.null_counts[column.ordinal - 1] > 0
            },
            "effective_as_of": self.effective_as_of.isoformat()
            if self.effective_as_of
            else None,
            "data_sha256": self.data_sha256,
            "schema_sha256": self.schema_sha256,
        }


def dataset_dir(value: str | Path | None = None) -> Path:
    """명시 인자 → 환경변수 → 로컬 공유 폴더 순으로 원천 위치를 정한다."""
    if value:
        return Path(value).expanduser().resolve()
    configured = os.environ.get("DATASET_DIR")
    if configured:
        return Path(configured).expanduser().resolve()
    return DEFAULT_DATASET_DIR.resolve()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_key_value(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_raw_value(value: object, data_type: str) -> object | None:
    """빈 값만 NULL로 바꾸고 숫자·코드의 0은 원본 그대로 보존한다."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    if isinstance(value, (date, datetime)):
        return value
    return value


def read_official_schema(
    path: Path, primary_key: tuple[str, ...] = ()
) -> tuple[OfficialColumn, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["schema"]:
            raise ValueError(f"{path.name}: schema 시트가 정확히 1개여야 합니다")
        rows = workbook["schema"].iter_rows(values_only=True)
        header = tuple(normalized_key_value(value) for value in next(rows))
        if header != SCHEMA_HEADER:
            raise ValueError(f"{path.name}: 공식 스키마 헤더 불일치 {header!r}")
        columns: list[OfficialColumn] = []
        for excel_row, row in enumerate(rows, start=2):
            if not any(normalized_key_value(value) for value in row):
                continue
            if len(row) != len(SCHEMA_HEADER):
                raise ValueError(f"{path.name}:{excel_row}: 스키마 행 폭 불일치")
            ordinal_text = normalized_key_value(row[0])
            name = normalized_key_value(row[1])
            data_type = normalized_key_value(row[2])
            nullable = normalized_key_value(row[3])
            description = normalized_key_value(row[4]) or ""
            if not ordinal_text or not name or not data_type or not nullable:
                raise ValueError(f"{path.name}:{excel_row}: 불완전한 스키마 행")
            if not TYPE_PATTERN.fullmatch(data_type):
                raise ValueError(
                    f"{path.name}:{excel_row}: 허용하지 않은 타입 {data_type!r}"
                )
            ordinal = int(ordinal_text)
            if ordinal != len(columns) + 1:
                raise ValueError(f"{path.name}:{excel_row}: 순번 {ordinal} 불연속")
            if nullable.upper() not in {"YES", "NO"}:
                raise ValueError(f"{path.name}:{excel_row}: Nullable={nullable!r}")
            columns.append(
                OfficialColumn(
                    ordinal=ordinal,
                    name=name.lower(),
                    data_type=data_type.lower(),
                    nullable=nullable.upper() == "YES",
                    description=description,
                )
            )
        names = [column.name for column in columns]
        if len(names) != len(set(names)):
            raise ValueError(f"{path.name}: 중복 컬럼명")
        return tuple(columns)
    finally:
        workbook.close()


def iter_data_rows(
    path: Path, columns: tuple[OfficialColumn, ...]
) -> Iterator[tuple[object | None, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["data"]:
            raise ValueError(f"{path.name}: data 시트가 정확히 1개여야 합니다")
        rows = workbook["data"].iter_rows(values_only=True)
        header = tuple(normalized_key_value(value) for value in next(rows))
        expected = tuple(column.name for column in columns)
        if header != expected:
            raise ValueError(f"{path.name}: data/schema 컬럼 1:1 불일치")
        for excel_row, row in enumerate(rows, start=2):
            if len(row) != len(columns):
                raise ValueError(
                    f"{path.name}:{excel_row}: 행 폭 {len(row)} != {len(columns)}"
                )
            yield tuple(
                normalize_raw_value(value, column.data_type)
                for value, column in zip(row, columns)
            )
    finally:
        workbook.close()


def parse_source_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalized_key_value(value)
    if not text:
        return None
    text = re.sub(r"\.0+$", "", text)
    if not re.fullmatch(r"\d{8}", text):
        return None
    try:
        return datetime.strptime(text, "%Y%m%d").date()
    except ValueError:
        return None


def exclusion_reason(
    spec: SourceSpec, row: tuple[object | None, ...], names: list[str]
) -> str | None:
    """260824 공식 정본은 행 단위 제외 없이 전량 보존한다."""
    return None


def inspect_source(root: Path, spec: SourceSpec) -> SourceInspection:
    data_path = root / spec.data_file
    schema_path = root / spec.schema_file
    columns = read_official_schema(schema_path, spec.primary_key)
    if len(columns) != spec.expected_columns:
        raise ValueError(
            f"{spec.code}: 스키마 {len(columns)}열 != 기대 {spec.expected_columns}열"
        )
    names = [column.name for column in columns]
    missing_pk = set(spec.primary_key) - set(names)
    if missing_pk:
        raise ValueError(f"{spec.code}: PK 컬럼 누락 {sorted(missing_pk)}")

    key_indexes = [names.index(name) for name in spec.primary_key]
    date_indexes = [names.index(name) for name in spec.effective_as_of_columns]
    seen: set[tuple[str, ...]] = set()
    maximum_as_of: date | None = None
    row_count = 0
    null_counts = [0] * len(columns)
    for excel_row, row in enumerate(iter_data_rows(data_path, columns), start=2):
        row_count += 1
        for index, value in enumerate(row):
            if value is None:
                null_counts[index] += 1
        key = tuple(normalized_key_value(row[index]) for index in key_indexes)
        if any(value is None for value in key):
            raise ValueError(f"{spec.code}:{excel_row}: PK NULL {key!r}")
        typed_key = tuple(value for value in key if value is not None)
        if typed_key in seen:
            raise ValueError(f"{spec.code}:{excel_row}: PK 중복 {typed_key!r}")
        seen.add(typed_key)
        for index in date_indexes:
            candidate = parse_source_date(row[index])
            if candidate and candidate > EXTERNAL_CUTOFF:
                raise LookAheadError(
                    f"{spec.code}:{excel_row}: {names[index]}={candidate}가 cutoff "
                    f"{EXTERNAL_CUTOFF} 초과"
                )
            if candidate:
                maximum_as_of = max(maximum_as_of or candidate, candidate)

    if row_count != spec.expected_rows:
        raise ValueError(
            f"{spec.code}: {row_count:,}행 != 기대 {spec.expected_rows:,}행"
        )
    return SourceInspection(
        spec=spec,
        data_path=data_path,
        schema_path=schema_path,
        columns=columns,
        null_counts=tuple(null_counts),
        row_count=row_count,
        effective_as_of=maximum_as_of,
        data_sha256=sha256_file(data_path),
        schema_sha256=sha256_file(schema_path),
    )


def validate_source_dir(value: str | Path | None = None) -> tuple[SourceInspection, ...]:
    root = dataset_dir(value)
    if not root.is_dir():
        raise FileNotFoundError(f"정본 데이터 디렉터리를 찾을 수 없습니다: {root}")
    expected = {
        name for spec in SOURCES for name in (spec.data_file, spec.schema_file)
    }
    candidates = [
        path
        for path in root.rglob("*.xlsx")
        if "__MACOSX" not in path.parts and not path.name.startswith("._")
    ]
    actual = {path.name for path in candidates}
    if len(candidates) != len(expected) or actual != expected:
        raise ValueError(
            "정상 XLSX 8개 집합 불일치: "
            f"누락={sorted(expected - actual)}, 초과={sorted(actual - expected)}, "
            f"정상파일수={len(candidates)}"
        )
    return tuple(inspect_source(root, spec) for spec in SOURCES)


def snapshot_hash(inspections: Iterable[SourceInspection]) -> str:
    payload = [
        {
            "file": item.data_path.name,
            "sha256": item.data_sha256,
            "schema_file": item.schema_path.name,
            "schema_sha256": item.schema_sha256,
        }
        for item in inspections
    ]
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()
